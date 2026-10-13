#!/usr/bin/env python3
"""Reads the attendee session allocation spreadsheet and writes tools/allocations.json,
which the delegate page uses for "Choose your name" (breakout workshop and Pillar Networking).

Usage (from the repository root):
    python3 tools/import_allocations.py "/path/to/Attendee Session Allocation.xlsx"

Only names, County FA, days attending and the two allocations are kept. The spreadsheet
itself stays out of this repository. The script also checks every County FA's Teamwork
and Collaboration and Lead, Manage, Do rooms against GROUPS in tools/content.py and stops
if they disagree, so the two pickers on the page can never contradict each other.
"""
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import content as C  # noqa: E402

WORKSHOPS = {"NLS Workshop": "National League System", "Data & Digital Workshop": "A Digital Journey"}
ROLE = {"Session Lead": "lead", "Facilitator": "facilitator"}


def split(value):
    """'Community, Main Room 1 (Facilitator)' -> ('Community, Main Room 1', 'facilitator')."""
    m = re.match(r"^(.*?)(?: \((Session Lead|Facilitator)\))?$", value.strip())
    return m.group(1), ROLE.get(m.group(2))


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    ws = openpyxl.load_workbook(sys.argv[1], data_only=True).worksheets[0]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

    people = {}
    problems = []
    groups = {c: g for g in C.GROUPS for c in g["counties"]}
    for r in rows:
        name = re.sub(r"\s*\(.*\)$", "", r["Full Name"].strip())
        entry = {"county": r.get("County FA") or "", "days": [n for n, k in ((1, "Day 1 Attending"), (2, "Day 2 Attending")) if r.get(k)]}
        for col, title in WORKSHOPS.items():
            v = r.get(col)
            if v and v not in ("Conference Host", "Conference Coordinator"):
                room, role = split(v)
                entry["workshop"] = {"title": title, "room": room, "role": role}
        v = r.get("Pillar Networking")
        if v:
            place, role = split(v)
            pillar, room = [x.strip() for x in place.split(",", 1)]
            entry["pillar"] = {"pillar": pillar, "room": room, "role": role}
        if "workshop" in entry or "pillar" in entry:
            people[name] = entry

        # County FA staff: their groups must match the county picker.
        county = r.get("County FA")
        if r.get("Organisation") == "County FA" and county:
            g = groups.get(county)
            if not g:
                problems.append(f"{county} ({name}) is not in GROUPS")
                continue
            for col, day, key in (("Teamwork & Collaboration", 1, "d1"), ("Lead-Manage-Do", 2, "d2")):
                v = r.get(col)
                if not v:
                    continue
                where, _ = split(v)
                want = f'{g[key]["room"]}, {g[key]["group"]}'
                if where != want:
                    problems.append(f"{name} ({county}) {col}: sheet says {where}, GROUPS says {want}")
    if problems:
        raise SystemExit("Fix these first:\n  " + "\n  ".join(problems))

    out = os.path.join(HERE, "allocations.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(dict(sorted(people.items())), f, ensure_ascii=False, indent=1)
    print(f"wrote {out}: {len(people)} people")


if __name__ == "__main__":
    main()
